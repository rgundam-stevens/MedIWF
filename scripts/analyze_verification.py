"""RQ4: human hand-check of the Tier-A detector (design-based precision / recall per rule, with stratum weights).

Inputs: the completed rater workbook: outputs/verification_sample_pass1_merged.xlsx (from merge_verification_sheets.py, the one-sheet-per-rule workbook folded flat) if present, else the _simple or plain pass-1 workbook; labels 0/1 and the sampling key
(outputs/verification_sample_key.csv, with stratum and stratum_population from make_verification_sample.py v3).
Each item carries the weight N_stratum / n_stratum, so the estimates refer to the whole corpus, not to the enriched sample.
Confidence intervals: stratified bootstrap (2,000 resamples within strata). Optionally a pass-2 workbook (a re-labelled
subset) gives intra-rater agreement per rule.
Usage: python3 scripts/analyze_verification.py [--pass2 outputs/verification_sample_pass2.xlsx] [outputs/analysis_verification_v1.txt]
"""
import sys, csv, pathlib, collections, io, argparse
import numpy as np
from openpyxl import load_workbook
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / 'src'))
from mediwf.stats import clopper_pearson, kappa_ci
root = pathlib.Path(__file__).resolve().parents[1]

# workbook column -> detector rule(s). meta_option covers both none/all of the above.
COLS = {"negative_stem": ["negative_stem"], "meta_option (all/none of the above)": ["none_of_the_above", "all_of_the_above"],
        "combination_options": ["combination_options"], "longest_option_key": ["longest_option_key"], "absolute_terms": ["absolute_terms"],
        "vague_terms": ["vague_terms"], "clang_cue": ["clang_cue"], "grammatical_cue": ["grammatical_cue"],
        "nonparallel_options": ["nonparallel_options"], "overlapping_options": ["overlapping_options"]}

def read_labels(path):
    ws = load_workbook(path).active
    hdr = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, row))
        if d.get("sample_id") is None: continue
        lab = {}
        for col in COLS:
            v = d.get(col)
            if v is None or str(v).strip() == "": lab[col] = None
            else: lab[col] = int(float(str(v).strip()))
        out[int(d["sample_id"])] = lab
    return out

def kappa(a, b):
    a = np.asarray(a); b = np.asarray(b); n = len(a)
    po = np.mean(a == b); pe = np.mean(a) * np.mean(b) + (1 - np.mean(a)) * (1 - np.mean(b))
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("out", nargs="?"); ap.add_argument("--pass1", default=str(next((root / f for f in ("outputs/verification_sample_pass1_merged.xlsx", "outputs/verification_sample_pass1_simple.xlsx") if (root / f).exists()), root / "outputs/verification_sample_pass1.xlsx")))
    ap.add_argument("--key", default=str(root / "outputs/verification_sample_key.csv")); ap.add_argument("--pass2", default=None); ap.add_argument("--boot", type=int, default=2000)
    a = ap.parse_args()
    buf = io.StringIO()
    def P(*x): print(*x); print(*x, file=buf)
    key = {int(r["sample_id"]): r for r in csv.DictReader(open(a.key))}
    labels = read_labels(a.pass1)
    n_strata = collections.Counter(r["stratum"] for r in key.values())
    weight = {sid: float(r["stratum_population"]) / n_strata[r["stratum"]] for sid, r in key.items()}
    ids = sorted(labels)
    missing = [sid for sid in ids if any(v is None for v in labels[sid].values())]
    P("items labelled: %d of %d; items with blank cells: %d" % (len(ids), len(key), len(missing)))
    P("strata (n sampled / population): " + ", ".join("%s %d/%s" % (st, n_strata[st], next(r["stratum_population"] for r in key.values() if r["stratum"] == st)) for st in sorted(n_strata)))
    P("note: rules with no detector-positive items in the corpus (negatives, meta-options, K-type) have no estimable precision; their recall rests on the unflagged stratum")
    rng = np.random.default_rng(20260916)
    strata = collections.defaultdict(list)
    for sid in ids: strata[key[sid]["stratum"]].append(sid)
    P("\n== Design-based precision and recall of each rule against the human labels (whole-corpus estimates; 95%% CI by stratified bootstrap, %d resamples)" % a.boot)
    P("%-38s %5s %5s %5s | %-22s %-22s | %s" % ("rule", "det+", "hum+", "both", "precision [CI]", "recall [CI]", "sample kappa"))
    exact = []
    for col, rules in COLS.items():
        det = {sid: int(any(key[sid][r] == "1" for r in rules)) for sid in ids}
        hum = {sid: labels[sid][col] for sid in ids}
        use = [sid for sid in ids if hum[sid] is not None]
        def est(sids):
            w = np.array([weight[s] for s in sids]); d = np.array([det[s] for s in sids]); h = np.array([hum[s] for s in sids])
            tp = (w * d * h).sum(); prec = tp / (w * d).sum() if (w * d).sum() else np.nan; rec = tp / (w * h).sum() if (w * h).sum() else np.nan
            return prec, rec
        prec, rec = est(use)
        boots = []
        for _ in range(a.boot):
            res = []
            for st, sids in strata.items():
                s2 = [s for s in sids if s in set(use)]
                if s2: res += list(rng.choice(s2, size=len(s2), replace=True))
            boots.append(est(res))
        bp = np.array([b[0] for b in boots], float); br = np.array([b[1] for b in boots], float)
        ci = lambda v: (np.nanpercentile(v, 2.5), np.nanpercentile(v, 97.5)) if np.isfinite(v).any() else (np.nan, np.nan)
        cp, cr = ci(bp), ci(br)
        dd = np.array([det[s] for s in use]); hh = np.array([hum[s] for s in use])
        P("%-38s %5d %5d %5d | %.2f [%.2f, %.2f]      %.2f [%.2f, %.2f]      | %.2f" % (col[:38], dd.sum(), hh.sum(), (dd & hh).sum(), prec, cp[0], cp[1], rec, cr[0], cr[1], kappa(dd, hh) if dd.sum() + hh.sum() else float("nan")))
        exact.append((col, int(dd.sum()), int(hh.sum()), int((dd & hh).sum()), int(((dd == 1) & (hh == 0)).sum()), int(((dd == 0) & (hh == 1)).sum()), int(((dd == 0) & (hh == 0)).sum())))
    P("\n== Unweighted sample estimates with exact (Clopper-Pearson) 95%% bounds and bootstrap kappa intervals (added 21 Sep 2026 at reviewers' request; rules with fewer than 10 positives are marked)")
    P("%-38s %-24s %-24s %s" % ("rule", "precision both/det+ [exact]", "recall both/hum+ [exact]", "kappa [bootstrap]"))
    for col, d, h, b, fp_, fn_, tn_ in exact:
        if d + h == 0: P("%-38s no positives" % col[:38]); continue
        cp1 = clopper_pearson(b, d) if d else (float("nan"), float("nan")); cp2 = clopper_pearson(b, h) if h else (float("nan"), float("nan")); kc = kappa_ci(b, fp_, fn_, tn_)
        flag = "  (<10 positives: interval uninformative)" if max(d, h) < 10 else ""
        P("%-38s %2d/%2d = %.2f [%.2f, %.2f]   %2d/%2d = %.2f [%.2f, %.2f]   %.2f [%.2f, %.2f]%s" % (col[:38], b, d, b / d if d else float("nan"), cp1[0], cp1[1], b, h, b / h if h else float("nan"), cp2[0], cp2[1], kappa(np.array([1]*b + [1]*fp_ + [0]*fn_ + [0]*tn_), np.array([1]*b + [0]*fp_ + [1]*fn_ + [0]*tn_)), kc[0], kc[1], flag))
    # composite: any of the ten hand-checked rules (detector) vs any human label; also without the clang rule
    P("\n== Composite over the ten hand-checked rules (design-weighted; whole-corpus estimates)")
    for label, cols in (("any of the 10 rules", list(COLS)), ("any of the 10 rules except clang", [c for c in COLS if c != "clang_cue"])):
        rules = [r for c in cols for r in COLS[c]]
        det = {sid: int(any(key[sid][r] == "1" for r in rules)) for sid in ids}
        hum = {sid: int(any(labels[sid][c] == 1 for c in cols)) for sid in ids if all(labels[sid][c] is not None for c in cols)}
        use = sorted(hum)
        def est2(sids):
            w = np.array([weight[s] for s in sids]); d = np.array([det[s] for s in sids]); h = np.array([hum[s] for s in sids])
            tp = (w * d * h).sum()
            return (tp / (w * d).sum() if (w * d).sum() else np.nan, tp / (w * h).sum() if (w * h).sum() else np.nan, (w * d).sum() / w.sum(), (w * h).sum() / w.sum())
        prec, rec, pd_, ph = est2(use)
        boots = []
        for _ in range(a.boot):
            res = []
            for st, sids in strata.items():
                s2 = [s for s in sids if s in hum]
                if s2: res += list(rng.choice(s2, size=len(s2), replace=True))
            boots.append(est2(res))
        b = np.array(boots, float)
        P("  %-34s precision %.2f [%.2f, %.2f]  recall %.2f [%.2f, %.2f]  | weighted prevalence: detector %.1f%%, human %.1f%% [%.1f, %.1f]" % (
            label, prec, np.nanpercentile(b[:, 0], 2.5), np.nanpercentile(b[:, 0], 97.5), rec, np.nanpercentile(b[:, 1], 2.5), np.nanpercentile(b[:, 1], 97.5),
            100 * pd_, 100 * ph, 100 * np.nanpercentile(b[:, 3], 2.5), 100 * np.nanpercentile(b[:, 3], 97.5)))
    if a.pass2:
        # v1.2 audit fix: pass-2 items carry NEW ids (make_pass2.py renumbers and shuffles them); map them back to pass-1 ids
        import json
        mapping = {int(m["pass2_id"]): int(m["pass1_id"]) for m in json.load(open(root / "outputs/verification_pass2_mapping_DO_NOT_OPEN.json"))}
        l2_raw = read_labels(a.pass2); l2 = {mapping[sid]: lab for sid, lab in l2_raw.items() if sid in mapping}
        both = [sid for sid in ids if sid in l2]
        P("\n== Intra-rater agreement, pass 1 vs pass 2 (%d items, matched through the id mapping)" % len(both))
        for col in COLS:
            x = [labels[s][col] for s in both if labels[s][col] is not None and l2[s][col] is not None]; y = [l2[s][col] for s in both if labels[s][col] is not None and l2[s][col] is not None]
            if x: P("  %-38s agreement %.3f  kappa %.2f  (positives pass1=%d pass2=%d)" % (col[:38], np.mean(np.array(x) == np.array(y)), kappa(x, y) if sum(x) + sum(y) else float("nan"), sum(x), sum(y)))
    if a.out: pathlib.Path(a.out).write_text(buf.getvalue()); print("\nwritten ->", a.out)

if __name__ == "__main__":
    main()
