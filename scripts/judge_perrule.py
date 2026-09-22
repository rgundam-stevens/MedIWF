"""Tier-B LLM judge, one call per (item, rule), using BenchMarker's per-rule prompt templates (MIT; see
config/prompts/benchmarker_rules/SOURCE.txt). Output format matches scripts/judge.py so validate_judge.py and
analyze_judge.py work unchanged. Resumable per item. Runs in Terminal (needs OpenRouter).

  python3 scripts/judge_perrule.py --judge openai/gpt-5.4-mini --validate --rules single_best_answer plausible_distractors no_extraneous_info focused_stem problem_in_stem
  python3 scripts/judge_perrule.py --judge openai/gpt-5.4-mini --corpus full --subsample 1600 --rules ...
"""
import argparse, json, ast, hashlib, pathlib, random, re, sys, time, datetime, collections
from concurrent.futures import ThreadPoolExecutor, as_completed
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
from mediwf.env import load_env, require
from mediwf.openrouter import OpenRouterClient

root = load_env()
TDIR = root / "config" / "prompts" / "benchmarker_rules"
ALL_RULES = ["single_best_answer", "plausible_distractors", "no_extraneous_info", "focused_stem", "problem_in_stem",
             "clear_language", "no_absolute_terms", "no_vague_terms", "no_logical_cues"]
SYSTEM = "You are an expert evaluator of multiple-choice questions. Follow the instructions exactly and answer only in JSON."
OUT = root / "data" / "generated"

def render(rule, item):
    tmpl = (TDIR / f"{rule}.txt").read_text()
    letters = sorted(item["options"])
    choices = str([item["options"][l] for l in letters])
    return tmpl.format(question=item["stem"], choices=choices, answer=item["answer"], lbrace="{", rbrace="}")

def parse_result(text):
    if not text: return None
    m = re.search(r"\{.*\}", text, re.S)
    for cand in ([text, m.group(0)] if m else [text]):
        try:
            d = json.loads(cand)
            r = str(d.get("result", "")).strip().lower()
            if r in ("pass", "fail"):
                return {"fail": r == "fail", "confidence": d.get("confidence"), "explanation": str(d.get("explanation", ""))[:300]}
        except Exception:
            continue
    return None

def corpus_items(run, subsample, seed=20260914):
    recs = {}
    for line in (OUT / f"items_{run}.jsonl").read_text().splitlines():
        r = json.loads(line)
        if r.get("parsed") and r["job_id"] not in recs: recs[r["job_id"]] = r
    items = [dict(uid=r["job_id"], model_label=r["model_label"], condition=r["condition"], specialty=r["specialty"], item=r["parsed"]) for r in recs.values()]
    if subsample:
        random.seed(seed); cells = collections.defaultdict(list)
        for it in items: cells[(it["model_label"], it["condition"])].append(it)
        per = subsample // len(cells); items = []
        for k in sorted(cells): items += random.sample(cells[k], min(per, len(cells[k])))
    return items

def validation_items():
    rows = [json.loads(l) for l in (root / "data" / "external" / "benchmarker" / "writing_flaws_judge.jsonl").read_text().splitlines()]
    seen = {}
    for r in rows:
        if not (isinstance(r["label"], int) or str(r["dataset"]).startswith("human_")): continue
        ch = r["choices"] if isinstance(r["choices"], list) else ast.literal_eval(r["choices"])
        key = hashlib.sha1((r["question"] + str(ch)).encode()).hexdigest()[:12]
        if key not in seen:
            seen[key] = dict(uid=key, dataset=r["dataset"], item={"stem": r["question"], "options": {chr(65 + i): c for i, c in enumerate(ch)}, "answer": r["answer"]})
    return list(seen.values())

def physician_items():
    """The 100 corpus items rated by the two physicians (outputs/content_rating/main_set_100.xlsx): uid = item id MAIN-001..100.
    Judged with the same per-rule prompts; scored against the physicians' labels by scripts/validate_judge_physician.py."""
    from openpyxl import load_workbook
    ws = load_workbook(root / "outputs" / "content_rating" / "main_set_100.xlsx", read_only=True)["main"]
    rows = list(ws.iter_rows(values_only=True)); out = []
    for r in rows[1:]:
        if not r[0]: continue
        opts = {L: str(v).strip() for L, v in zip("ABCDE", r[2:7]) if v not in (None, "")}
        out.append(dict(uid=str(r[0]).strip(), dataset="physician100", item={"stem": str(r[1]), "options": opts, "answer": str(r[7]).strip()}))
    assert len(out) == 100, len(out)
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--judge", required=True); ap.add_argument("--corpus"); ap.add_argument("--subsample", type=int, default=0)
    ap.add_argument("--validate", action="store_true"); ap.add_argument("--physician", action="store_true", help="judge the 100 physician-rated corpus items (Amendment 6)")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--rules", nargs="*", default=ALL_RULES)
    ap.add_argument("--limit", type=int, default=0, help="pilot: judge only the first N items (deterministic order)")
    ap.add_argument("--max-cost", type=float, default=0.0, help="stop submitting new items once the cumulative cost of this run reaches this many USD (in-flight items finish; resumable)")
    ap.add_argument("--reasoning", action="store_true", help="Amendment 6 A28: enable hidden reasoning (effort high unless --reasoning-budget is given); output file gets the suffix _reasoning; a verdict is kept only if the response reports reasoning tokens")
    ap.add_argument("--reasoning-budget", type=int, default=0, help="optional thinking-token budget per call (0 = provider default at effort high)")
    a = ap.parse_args()
    rules = [r for r in a.rules if r in ALL_RULES]
    if a.validate: items, target = validation_items(), "benchmarker"
    elif a.physician: items, target = physician_items(), "physician100"
    elif a.corpus: items, target = corpus_items(a.corpus, a.subsample), (a.corpus + (f"_sub{a.subsample}" if a.subsample else ""))
    else: ap.error("--corpus RUN, --validate or --physician")
    n_all = len(items)
    if a.limit: items = items[:a.limit]
    slug = re.sub(r"[^a-z0-9]+", "-", a.judge.lower())
    if a.reasoning: target += "_reasoning" + (f"_budget{a.reasoning_budget}" if a.reasoning_budget else "")
    out = OUT / f"judgeR_{slug}_{target}.jsonl"
    reasoning_req = ({"max_tokens": a.reasoning_budget} if a.reasoning_budget else {"effort": "high"}) if a.reasoning else None
    done = {}
    if out.exists():
        for line in out.read_text().splitlines():
            try:
                rec = json.loads(line)
                if rec.get("judgment"): done[rec["uid"]] = rec
            except Exception: pass
    # an item is complete when every requested rule has a verdict
    todo = [it for it in items if not (it["uid"] in done and all(r in done[it["uid"]]["judgment"] for r in rules))]
    print(f"judge={a.judge} target={target} rules={len(rules)}: {len(items)} items, {len(items)-len(todo)} complete, {len(todo)} to run ({len(todo)*len(rules)} calls)")
    client = OpenRouterClient(require("OPENROUTER_API_KEY"))
    def work(it):
        prev = done.get(it["uid"], {}).get("judgment", {}) or {}
        details = done.get(it["uid"], {}).get("details", {}) or {}
        cost = 0.0; ptok = ctok = rtok = 0; errors = {}; meta = {}
        for rule in rules:
            if rule in prev: continue
            res = client.chat(a.judge, SYSTEM, render(rule, it["item"]), temperature=0.0, seed=7, max_tokens=(16000 if a.reasoning else 600), reasoning=reasoning_req)
            cost += float(res.get("cost_usd") or 0); ptok += int(res.get("prompt_tokens") or 0); ctok += int(res.get("completion_tokens") or 0); rtok += int(res.get("reasoning_tokens") or 0)
            if a.reasoning: meta[rule] = {"reasoning_tokens": res.get("reasoning_tokens"), "temperature_sent": res.get("temperature_sent"), "dropped": res.get("dropped_params"), "finish": res.get("finish_reason")}
            pr = parse_result(res.get("content")) if res.get("ok") else None
            if pr is not None and a.reasoning and res.get("reasoning_tokens") is None:
                # the usage block did not report reasoning at all: the setting was not honoured, so the verdict is not kept.
                # A reported count of 0 means the model was allowed to think and chose not to (adaptive thinking); it is kept and counted.
                errors[rule] = "reasoning requested but the response reported no reasoning field (finish=%s); verdict discarded" % res.get("finish_reason"); pr = None
            if pr is None:
                errors.setdefault(rule, (res.get("error") or (res.get("content") or "")[:200]))
            else:
                prev[rule] = pr["fail"]; details[rule] = {"confidence": pr["confidence"], "explanation": pr["explanation"]}
        rec = {k: v for k, v in it.items() if k != "item"}
        rec.update(judge=a.judge, mode="perrule", requested_at=datetime.datetime.utcnow().isoformat() + "Z",
                   judgment=prev if prev else None, details=details, errors=errors, cost_usd=cost, prompt_tokens=ptok, completion_tokens=ctok,
                   reasoning_tokens=rtok, reasoning=(dict(reasoning_req, meta=meta) if a.reasoning else None))
        return rec
    cost = 0.0; ok = 0; t0 = time.time(); done_n = 0; stopped = False
    with out.open("a") as f, ThreadPoolExecutor(max_workers=a.workers) as ex:
        pending = set(); queue = list(todo)
        while queue or pending:
            # submit up to `workers` items, but not once the cost cap is reached
            while queue and len(pending) < a.workers and not stopped:
                pending.add(ex.submit(work, queue.pop(0)))
            if not pending: break
            fut = next(as_completed(pending)); pending.discard(fut)
            rec = fut.result(); f.write(json.dumps(rec, ensure_ascii=False) + "\n"); f.flush(); done_n += 1
            cost += float(rec.get("cost_usd") or 0); ok += bool(rec.get("judgment")) and not rec["errors"]
            if done_n % 25 == 0 or done_n == len(todo): print(f"  {done_n}/{len(todo)} items | complete={ok} | cost ${cost:.3f} | {time.time()-t0:.0f}s")
            if a.max_cost and cost >= a.max_cost and not stopped:
                stopped = True; print(f"  cost cap ${a.max_cost:.2f} reached after {done_n} items; no new items submitted (re-run the same command to resume)")
    if done_n:
        per_item = cost / done_n
        print(f"this run: {done_n} items, ${cost:.2f} (${per_item:.3f} per item); projected cost for all {n_all} items: ${per_item * n_all:.2f}")
        if a.reasoning:
            recs = [json.loads(l) for l in out.read_text().splitlines() if l.strip()]
            metas = [m for r in recs for m in ((r.get("reasoning") or {}).get("meta") or {}).values()]
            thought = [m["reasoning_tokens"] for m in metas if m.get("reasoning_tokens")]
            print(f"reasoning check: {len(metas)} calls; the model used thinking on {len(thought)} of them ({100 * len(thought) / max(len(metas), 1):.0f}%), median {sorted(thought)[len(thought) // 2] if thought else 0} thinking tokens when used; calls with no reasoning field at all: {sum(1 for m in metas if m.get('reasoning_tokens') is None)}")
    print("finished ->", out)

if __name__ == "__main__":
    main()
