"""Build the physician content-rating package (Amendment 2): a 20-item calibration set with planted defects, the
100-item main set, a 2-item live test, and the hidden key. Rater sheets contain stem, options and key only (no
rationale, no model). Usage: python3 scripts/make_content_rating_sheets.py"""
import json, csv, random, pathlib, copy
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
root = pathlib.Path(__file__).resolve().parents[1]
random.seed(20260916)
sel = json.load(open(root / "outputs" / "content_rating_selection.json"))
recs = {}
for line in (root / "data" / "generated" / "items_full.jsonl").read_text().splitlines():
    r = json.loads(line)
    if r.get("parsed") and r["job_id"] not in recs: recs[r["job_id"]] = r
cal = sel["calibration_candidates"]   # C1..C26 in order
def C(n): return cal[n - 1]

# ---- calibration design: 10 unmodified (expected clean), 2 natural defects, 8 planted defects
plan = []   # (job_id, kind, modify_fn, expected)
def keep(n, kind="clean", expected=None, note=""):
    plan.append((C(n), kind, None, expected or {}, note))
def planted(n, fn, expected, note):
    plan.append((C(n), "planted", fn, expected, note))

for n in (1, 2, 4, 7, 10, 13, 15, 17, 20, 21): keep(n)
keep(26, "natural", {"q1": "no/unsure"}, "Levofloxacin or moxifloxacin monotherapy is also guideline-acceptable for non-severe CAP: more than one defensible answer")
keep(24, "natural", {"q1": "no/unsure"}, "Watchful waiting is an accepted option for an asymptomatic inguinal hernia; rationale contradicts the vignette (says symptomatic)")

def set_key(letter):
    def f(p): p["answer"] = letter
    return f
def set_option(letter, text):
    def f(p): p["options"][letter] = text
    return f
def stem_replace(old, new):
    def f(p):
        assert old in p["stem"], old; p["stem"] = p["stem"].replace(old, new)
    return f
def chain(*fs):
    def f(p):
        for g in fs: g(p)
    return f

planted(3, set_key("E"), {"q1": "no"}, "Wrong key: hypovolaemic, ketotic patient needs IV fluids first, not doxylamine-pyridoxine")
planted(14, set_key("E"), {"q1": "no"}, "Wrong key: jaundice <24 h with ABO setup needs a direct antiglobulin test, not TSH")
planted(19, set_option("D", "Intravenous cefotaxime and intravenous albumin"), {"q1": "no/unsure"}, "Two defensible answers: cefotaxime + albumin is equivalent to ceftriaxone + albumin")
planted(9, set_option("B", "Amphetamine-dextroamphetamine"), {"q1": "no/unsure"}, "Two defensible answers: both stimulants are first-line initial pharmacotherapy")
planted(11, stem_replace("Hemoglobin is 9.8 g/dL", "Hemoglobin is 15.8 g/dL"), {"q3": "no"}, "Internally inconsistent: normal haemoglobin but the item is built around iron-deficiency anaemia")
planted(18, stem_replace("37 weeks and 2 days gestation", "31 weeks and 2 days gestation"), {"q1": "no", "q3": "no"}, "At 31 weeks without severe features, induction is wrong (expectant management + steroids); key no longer fits the vignette")
planted(8, set_option("E", "Amputation of the right lower extremity"), {"q2": "no"}, "Implausible distractor: a surgical procedure among histologic patterns")
planted(25, set_option("B", "Immediate initiation of chemotherapy"), {"q2": "no"}, "Implausible distractor: chemotherapy without a diagnosis")

items = []
for jid, kind, fn, expected, note in plan:
    p = copy.deepcopy(recs[jid]["parsed"])
    if fn: fn(p)
    items.append(dict(job_id=jid, kind=kind, expected=expected, note=note, item=p, model=recs[jid]["model_label"]))
random.shuffle(items)

HDR = ["item_id", "stem", "A", "B", "C", "D", "E", "keyed_answer",
       "Q1 keyed answer is the single best answer (yes/no/unsure)", "Q2 every distractor plausible (yes/no)",
       "Q3 vignette accurate and consistent (yes/no)", "Q4 tests a meaningful decision (yes/no)",
       "Q5 accept as written (yes/minor edits/no)", "justification (required for every no/unsure)", "minutes spent"]
def write_sheet(rows, path, sheetname):
    wb = Workbook(); ws = wb.active; ws.title = sheetname
    ws.append(HDR)
    for c in ws[1]: c.font = Font(bold=True, name="Arial", size=10); c.fill = PatternFill("solid", fgColor="DDEBF7"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for rid, p in rows:
        o = p["options"]
        ws.append([rid, p["stem"], o.get("A"), o.get("B"), o.get("C"), o.get("D"), o.get("E"), p["answer"]] + [""] * 7)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10)
    for col, w in zip("ABCDEFGH", [9, 75, 24, 24, 24, 24, 24, 8]): ws.column_dimensions[col].width = w
    for col in "IJKLMNO": ws.column_dimensions[col].width = 16
    ws.freeze_panes = "C2"; ws.row_dimensions[1].height = 60
    wb.save(path)

out = root / "outputs" / "content_rating"; out.mkdir(exist_ok=True)
write_sheet([("CAL-%02d" % i, it["item"]) for i, it in enumerate(items, 1)], out / "calibration_set_20.xlsx", "calibration")
with open(out / "calibration_key_DO_NOT_SEND.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["item_id", "job_id", "model", "kind", "expected", "note"])
    for i, it in enumerate(items, 1): w.writerow(["CAL-%02d" % i, it["job_id"], it["model"], it["kind"], json.dumps(it["expected"]), it["note"]])

main = sel["main"]
write_sheet([("MAIN-%03d" % i, recs[j]["parsed"]) for i, j in enumerate(main, 1)], out / "main_set_100.xlsx", "main")
with open(out / "main_key_DO_NOT_SEND.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["item_id", "job_id", "model", "condition", "specialty"])
    for i, j in enumerate(main, 1): w.writerow(["MAIN-%03d" % i, j, recs[j]["model_label"], recs[j]["condition"], recs[j]["specialty"]])

# live test for the platform chat: one clean item (C22) and one with a wrong key (C12 keyed to annual ultrasonography)
lt = []
p = copy.deepcopy(recs[C(22)]["parsed"]); lt.append(("LIVE-1 (expected: all yes)", p))
p = copy.deepcopy(recs[C(12)]["parsed"]); p["answer"] = "B"; lt.append(("LIVE-2 (expected: Q1 no - one-time screening, not annual)", p))
with open(out / "live_test_for_chat.txt", "w") as f:
    for label, p in lt:
        f.write(label + "\n" + p["stem"] + "\n")
        for L in "ABCDE": f.write("  %s. %s\n" % (L, p["options"][L]))
        f.write("  Keyed answer: %s\n\n" % p["answer"])
print("written to", out); print("calibration order:", [(("CAL-%02d" % i), it["kind"]) for i, it in enumerate(items, 1)])
