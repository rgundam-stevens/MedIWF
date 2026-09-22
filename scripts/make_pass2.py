"""Prepare the pass-2 workbook for the hand-check: 150 items of pass 1 (a stratified random subset, all strata represented),
in a new random order and with new sample ids, so that the second reading is blind to the first.
Usage: python3 scripts/make_pass2.py [--n 150]   (run on or after the date agreed in the protocol, >= 4 weeks after pass 1)
"""
import csv, random, pathlib, argparse, collections, json
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill
root = pathlib.Path(__file__).resolve().parents[1]
ap = argparse.ArgumentParser(); ap.add_argument("--n", type=int, default=150); ap.add_argument("--seed", type=int, default=20261014); a = ap.parse_args()
random.seed(a.seed)
key = list(csv.DictReader(open(root / "outputs/verification_sample_key.csv")))
by = collections.defaultdict(list)
for r in key: by[r["stratum"]].append(int(r["sample_id"]))
share = a.n / len(key); pick = []
for st, sids in by.items(): pick += random.sample(sids, max(1, round(share * len(sids))))
rest = [int(r["sample_id"]) for r in key if int(r["sample_id"]) not in set(pick)]
random.shuffle(rest); pick = (pick + rest)[:a.n]; random.shuffle(pick)   # v1.2: exactly n items (rounding gave 149)
src = root / "outputs/verification_sample_pass1_simple.xlsx"
if not src.exists(): src = root / "outputs/verification_sample_pass1.xlsx"
ws1 = load_workbook(src)["pass1"] if "pass1" in load_workbook(src).sheetnames else load_workbook(src).active
rows = {int(r[0]): r for r in ws1.iter_rows(min_row=2, values_only=True) if r[0] is not None}
hdr = [c.value for c in ws1[1]]
wb = Workbook(); ws = wb.active; ws.title = "pass2"; ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True, name="Arial"); c.fill = PatternFill("solid", fgColor="FCE4D6")
mapping = []
for new_id, old_id in enumerate(pick, 1):
    r = list(rows[old_id]); r[0] = new_id
    for k in range(8, 19): r[k] = None               # blank the ten labels and the notes; helper columns (if any) are kept
    ws.append(r); mapping.append({"pass2_id": new_id, "pass1_id": old_id})
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10)
for col, w in zip("ABCDEFGH", [9, 70, 22, 22, 22, 22, 22, 6]): ws.column_dimensions[col].width = w
from openpyxl.utils import get_column_letter
for j in range(20, ws.max_column + 1): ws.column_dimensions[get_column_letter(j)].width = 42
ws.freeze_panes = "C2"
wb.save(root / "outputs/verification_sample_pass2.xlsx")
json.dump(mapping, open(root / "outputs/verification_pass2_mapping_DO_NOT_OPEN.json", "w"))
import sys; sys.path.insert(0, str(root / "scripts")); from make_verification_helpers import build_byrule
items = [{"sample_id": new_id, "stem": rows[old_id][1], **{L: rows[old_id][2 + i] for i, L in enumerate("ABCDE")}, "key": str(rows[old_id][7] or "").strip().upper()[:1]} for new_id, old_id in enumerate(pick, 1)]
build_byrule(items, root / "outputs/verification_sample_pass2_byrule.xlsx", "pass 2")
print("wrote outputs/verification_sample_pass2.xlsx (%d items) and the id mapping (do not open until pass 2 is done)" % len(pick))
