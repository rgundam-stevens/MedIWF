"""Build the simplified hand-check workbook: the same 439 items as outputs/verification_sample_pass1.xlsx, the same ten
label columns (so analyze_verification.py reads it unchanged), plus grey "helper" columns that do the counting for the
rater: option lengths, the Rule 4 arithmetic, number-first and full-sentence options, the word-count spread, the listed
sweeping and frequency words that occur, the Rule 7 candidate words, the Rule 10 containment pairs and the a/an ending.
The helpers are computed here with plain string operations, independently of the detector (src/mediwf/detector.py is
not imported), and they never show the detector's flags. Two extra sheets carry a one-page routine and a cheat sheet.
Usage: python3 scripts/make_verification_helpers.py   ->  outputs/verification_sample_pass1_simple.xlsx (flat, one row per item)
                                                        and outputs/verification_sample_pass1_byrule.xlsx (one sheet per rule: the rater's copy)
The by-rule workbook shows, on each sheet, only what that rule needs (the key letter appears only on the sheets for rules 4, 7 and 8) and one
answer column; scripts/merge_verification_sheets.py folds the ten answer columns back into the flat layout that analyze_verification.py reads.
"""
import re, pathlib, statistics
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
root = pathlib.Path(__file__).resolve().parents[1]
SRC = root / "outputs/verification_sample_pass1.xlsx"
OUT = root / "outputs/verification_sample_pass1_simple.xlsx"

ABSOLUTE = ["always", "never", "all", "none", "only", "every", "completely", "entirely", "absolutely", "totally", "definitely", "impossible", "invariably", "exclusively"]
VAGUE = ["usually", "often", "frequently", "sometimes", "rarely", "seldom", "occasionally", "generally", "regularly", "infrequently", "commonly"]
FIXED = [r"every (\d+|other|day|night|morning|evening|few)", r"once daily", r"only (child|one|two|three|a few|\d+)", r"all-cause", r"at all", r"(all|none) of the above", r"all of these"]
GENERIC = set("patient patients treatment management therapy diagnosis cause test study appropriate initial next step likely most following "
              "with from that this which than after before during without into over under between should would could been being have "
              "their there these those where when what more less other some same such also because".split())
LABELS = ["negative_stem", "meta_option (all/none of the above)", "combination_options", "longest_option_key", "absolute_terms", "vague_terms",
          "clang_cue", "grammatical_cue", "nonparallel_options", "overlapping_options"]


def words(s): return [w for w in re.split(r"\s+", (s or "").strip()) if w]
def clean_word(w): return re.sub(r"^[^A-Za-z0-9]+|[^A-Za-z0-9]+$", "", w).lower()
def base(w):
    """Singular/plural folding only (the rater guide treats singular and plural as the same word): -ies -> -y,
    -xes/-ches/-shes/-sses/-zes -> drop es, other -s -> drop s (not -ss). No -ing/-ed stripping."""
    w = clean_word(w)
    if len(w) > 4 and w.endswith("ies"): return w[:-3] + "y"
    if len(w) > 4 and w.endswith(("xes", "zes", "ches", "shes", "sses")): return w[:-2]
    if len(w) > 4 and w.endswith("s") and not w.endswith("ss"): return w[:-1]
    return w
def is_sentence(s): return len(words(s)) >= 4 and (s or "").strip().endswith(".")
def starts_number(s): return bool(re.match(r"^\s*[\d.]", s or ""))


def find_listed(options, wordlist):
    hits = []
    for L, o in options.items():
        for w in wordlist:
            for m in re.finditer(r"\b" + w + r"\b", o or "", flags=re.I):
                ctx = (o or "")[max(0, m.start() - 25): m.end() + 25].replace("\n", " ")
                fixed = any(re.search(f, (o or "")[max(0, m.start() - 12): m.end() + 14], flags=re.I) for f in FIXED)
                hits.append("%s (%s: …%s…)%s" % (w, L, ctx.strip(), "  ← looks like a fixed phrase, check the exemption" if fixed else ""))
    return "; ".join(hits)


def clang_candidates(stem, options, key):
    stem_bases = {base(w) for w in words(stem)}
    out = []
    seen = set()
    for w in words(options.get(key, "")):
        b = base(w); c = clean_word(w)
        if len(c) < 4 or c in GENERIC or b in seen or not re.search(r"[a-z]", c): continue
        seen.add(b)
        if b not in stem_bases: continue
        also = [L for L, o in options.items() if L != key and b in {base(x) for x in words(o)}]
        out.append("%s → in the stem, %s" % (c, ("also in wrong option(s) " + ", ".join(also)) if also else "in NO wrong option  ← candidate"))
    return "; ".join(out) if out else "no key word (4+ letters, non-generic) is in the stem"


def overlap_pairs(options):
    norm = {L: re.sub(r"\s+", " ", (o or "").strip().rstrip(".").lower()) for L, o in options.items()}
    pairs = []
    for a, x in norm.items():
        if len(x) < 6: continue
        for b, y in norm.items():
            if a != b and re.search(r"(?<![A-Za-z0-9])" + re.escape(x) + r"(?![A-Za-z0-9])", y): pairs.append("%s inside %s" % (a, b))
    return "; ".join(pairs) if pairs else "none found"


def question_ending(stem):
    s = (stem or "").strip()
    # last sentence: text after the last . ! ? that precedes the end
    parts = re.split(r"(?<=[.!?])\s+", s)
    last = parts[-1] if parts else s
    last = last.rstrip("?:. ").strip()
    lw = clean_word(last.split()[-1]) if last.split() else ""
    return ("question sentence ends with '%s'  ← do the a/an check" % lw) if lw in ("a", "an") else "does not end with a/an"


def helpers_for(stem, options, key):
    ch = {L: len((o or "").strip()) for L, o in options.items()}
    wc = {L: len(words(o)) for L, o in options.items()}
    wrong = [L for L in options if L != key]
    longest_wrong = max(wrong, key=lambda L: ch[L]) if wrong else ""
    need = 1.25 * ch.get(longest_wrong, 0)
    lengths = " · ".join("%s %d ch / %d w" % (L, ch[L], wc[L]) for L in options)
    r4 = "key %s: %d characters, %d words; longest wrong option %s: %d characters; 1.25 × %d = %.2f → mark 1 only if key has 4+ words AND more than %.2f characters" % (
        key, ch.get(key, 0), wc.get(key, 0), longest_wrong, ch.get(longest_wrong, 0), ch.get(longest_wrong, 0), need, need)
    nums = [L for L, o in options.items() if starts_number(o)]
    sents = [L for L, o in options.items() if is_sentence(o)]
    counts = [wc[L] for L in options]
    mean = statistics.mean(counts) if counts else 0
    cv = (statistics.pstdev(counts) / mean) if mean else 0
    r9 = "options starting with a number: %s; full sentences (4+ words, ends with a full stop): %s; word counts %s → spread ÷ average = %.2f (average %.1f)%s" % (
        ", ".join(nums) if nums else "none", ", ".join(sents) if sents else "none", "/".join(str(c) for c in counts), cv, mean,
        "  ← above 0.6 with average ≥ 3" if (cv > 0.6 and mean >= 3) else "")
    r8 = "%s; key is a full sentence: %s; wrong options that are full sentences: %s" % (
        question_ending(stem), "yes" if key in sents else "no", ", ".join(L for L in sents if L != key) or "none")
    return {
        "option lengths (characters / words)": lengths,
        "Rule 4 arithmetic": r4,
        "Rule 5 sweeping words found": find_listed(options, ABSOLUTE) or "none of the listed words",
        "Rule 6 frequency words found": find_listed(options, VAGUE) or "none of the listed words",
        "Rule 7 candidate words (key words that are in the stem)": clang_candidates(stem, options, key),
        "Rule 8 helpers": r8,
        "Rule 9 helpers": r9,
        "Rule 10 containment pairs": overlap_pairs(options),
    }


HOWTO = [
    ("How to do the hand-check, in plain words", True),
    ("", False),
    ("You are proof-reading the FORM of 439 questions, not their medicine. For every row you answer ten yes/no questions and type 1 (yes) or 0 (no) in the ten blue columns. Never leave a blue cell blank.", False),
    ("", False),
    ("The grey columns on the right do the counting for you. They are computed from the text of the item, not by the checker you are testing, and they never tell you the answer to a judgement call. Use them, but the decision and the 1/0 are yours.", False),
    ("", False),
    ("The routine for each row", True),
    ("1. Read the last sentence of the stem (the question). Column negative_stem: 1 if it asks for the wrong, false, least or excluded option (NOT / EXCEPT / LEAST / FALSE / INCORRECT in that sentence).", False),
    ("2. Read the five options as a list. Column meta_option: any 'all / none / neither of the above'? Column combination_options: options built from other options ('A and B', 'I and III')? Column absolute_terms: look at the grey 'Rule 5' cell; mark 1 if a listed word is there and it is NOT a fixed phrase (every 8 hours, only child, all-cause...). Column vague_terms: grey 'Rule 6' cell; mark 1 if any listed word is there. Column nonparallel_options: grey 'Rule 9' cell; mark 1 if number-first options are mixed with text options, or full sentences are mixed with fragments, or the spread ratio is above 0.6 with average 3 or more. Column overlapping_options: grey 'Rule 10' cell; mark 1 if one option's whole text sits inside another (confirm it is whole words).", False),
    ("3. Now look at the key. Column longest_option_key: grey 'Rule 4' cell gives the numbers; mark 1 only if the key has 4+ words AND more than 1.25 times the characters of the longest wrong option. Column clang_cue: grey 'Rule 7' cell lists key words that are also in the stem; mark 1 if at least one of them is a DISTINCTIVE word (drug, disease, body part, finding, procedure) and is in no wrong option. Column grammatical_cue: grey 'Rule 8' cell; mark 1 if the question ends in a/an and only some options fit it, OR the key alone is a full sentence while no wrong option is.", False),
    ("", False),
    ("Notes column: write 'category mix' (diagnoses next to treatments), 'logical overlap' (less than 5 mg / less than 10 mg), or anything odd (six options, no key). Type the date at the start of each sitting.", False),
    ("", False),
    ("Practical rules", True),
    ("Do the ten practice items in the Rater Guide (docs/MedIWF_Rater_Guide_v3.docx) first and check them against its answer key. Work in sittings of about 40 rows. Do not look at any detector output, do not ask an AI, and do not go back to change earlier answers because a later item made you rethink; note the doubt instead. When you have finished all 439 rows, tell me; the second reading of 150 rows happens at least four weeks later, from a fresh workbook.", False),
]

CHEAT = [
    ("Column", "Ask yourself", "Mark 1 when"),
    ("negative_stem", "Does the question sentence ask for the wrong one?", "NOT, EXCEPT, LEAST, FALSE or INCORRECT in the question sentence makes the examinee pick the option that does not apply"),
    ("meta_option", "Any 'all / none / neither of the above'?", "Any option is one of these, whether or not it is the key"),
    ("combination_options", "Options made from other options?", "Any option refers to two or more other options by letter or numeral. A lone numeral (stage IV) and ordinary two-drug regimens do not count"),
    ("longest_option_key", "Is the key a quarter longer than every distractor?", "Key has at least 4 words and MORE than 1.25 x the characters of the longest wrong option (grey Rule 4 cell)"),
    ("absolute_terms", "Sweeping words in any option?", "always, never, all, none, only, every, completely, entirely, absolutely, totally, definitely, impossible, invariably, exclusively - except fixed phrases (every 8 hours, only child, all-cause, at all, only one/two/a few) and the 'all/none of the above' options that belong to meta_option"),
    ("vague_terms", "Frequency words in any option?", "usually, often, frequently, sometimes, rarely, seldom, occasionally, generally, regularly, infrequently, commonly (options only, not the stem)"),
    ("clang_cue", "A specific stem word only in the key?", "A distinctive word (drug, disease, body part, finding, procedure) is in the stem, in the key, and in none of the wrong options; singular and plural count as the same word"),
    ("grammatical_cue", "Does grammar or form single out the key?", "Question ends in a/an and only some options fit it; or the key alone is a full sentence (4+ words, final full stop) while no wrong option is"),
    ("nonparallel_options", "Same shape?", "Number-first options mixed with text options; full sentences mixed with fragments; word-count spread / average above 0.6 with average 3 or more. Mixed categories go in the notes column, not here"),
    ("overlapping_options", "One option inside another?", "One option's whole text (6+ characters) appears inside another option as whole words, ignoring capitals and a final full stop. Logical overlap without shared text goes in the notes column"),
]


BYRULE = [  # (sheet name, label column in the flat layout, which fields to show, helper column to show, the one-line rule)
    ("1 negative_stem", "negative_stem", ["stem"], None,
     "Read the whole stem; the question is normally its last sentence. Mark 1 if that question asks for the wrong, false, least or excluded option (NOT, EXCEPT, LEAST, FALSE, INCORRECT, 'should be avoided'). A 'not' in the patient story or inside an option is 0; a question asking for the best, most likely or next option is 0."),
    ("2 meta_option", "meta_option (all/none of the above)", ["options"], None,
     "Mark 1 if any option is 'all of the above', 'none of the above', 'neither of the above' or 'all of these', key or not. 'No treatment is needed' is a real answer: 0."),
    ("3 combination", "combination_options", ["options"], None,
     "Mark 1 if any option refers to two or more OTHER options by letter or numeral ('A and B', 'both A and C', 'I, II and III', '1 and 3'). A lone numeral (stage IV) is 0; ordinary two-drug regimens ('aspirin and clopidogrel') are 0."),
    ("4 longest", "longest_option_key", ["options", "key"], "Rule 4 arithmetic",
     "Mark 1 only if the KEY has at least 4 words AND more than 1.25 x the characters of the longest WRONG option. The grey cell gives the numbers; a long wrong option never counts here."),
    ("5 absolute", "absolute_terms", ["options"], "Rule 5 sweeping words found",
     "Mark 1 if any option contains a listed sweeping word (always, never, all, none, only, every, completely, entirely, absolutely, totally, definitely, impossible, invariably, exclusively) UNLESS it is a fixed phrase: every 8 hours, every other day, once daily, only child, all-cause, at all, only one/two/a few, or an 'all/none of the above' option (that is sheet 2). The grey cell lists what was found; you judge the exemption."),
    ("6 vague", "vague_terms", ["options"], "Rule 6 frequency words found",
     "Mark 1 if any option contains a listed frequency word: usually, often, frequently, sometimes, rarely, seldom, occasionally, generally, regularly, infrequently, commonly. 'Regular' and 'common' are not on the list: 0. The stem is never checked."),
    ("7 clang", "clang_cue", ["stem", "options", "key"], "Rule 7 candidate words (key words that are in the stem)",
     "Mark 1 if a DISTINCTIVE word (drug, disease, organism, body part, finding, procedure, mechanism; 4+ letters) is in the stem, in the key, and in NO wrong option. Generic words (patient, treatment, management, therapy, diagnosis, cause, test, next, step, likely...) never count. Singular and plural are the same word. The grey cell lists the candidates; you judge whether the word is distinctive."),
    ("8 grammatical", "grammatical_cue", ["stem", "options", "key"], "Rule 8 helpers",
     "Two checks, either gives a 1. (a) The question sentence ends in 'a' or 'an' and only SOME options fit it (an + vowel sound, a + consonant sound). (b) The KEY alone is a full sentence (4+ words, ends with a full stop) while no wrong option is. An -ing word or a plural at the start of the key is NOT a cue."),
    ("9 nonparallel", "nonparallel_options", ["options"], "Rule 9 helpers",
     "Mark 1 if (a) two or more options start with a number while at least one does not, or (b) some options are full sentences (4+ words, final full stop) and others are fragments, or (c) the word-count spread / average is above 0.6 with an average of at least 3 (the grey cell computes it). Mixed CATEGORIES (a diagnosis next to treatments) are 0 here: write 'category mix' in notes."),
    ("10 overlapping", "overlapping_options", ["options"], "Rule 10 containment pairs",
     "Mark 1 if one option's WHOLE text (6+ characters) appears inside another option as whole words, ignoring capitals and a final full stop ('Aspirin' inside 'Aspirin and clopidogrel'). Part of a word ('800 mL' in '4800 mL') is 0. Logical overlap without shared text ('less than 5 mg' / 'less than 10 mg') is 0: write 'logical overlap' in notes."),
]


def question_sentence(stem):
    s = (stem or "").strip()
    parts = re.split(r"(?<=[.!?])\s+", s)
    return parts[-1] if parts else s


def build_byrule(items, out_path, pass_name):
    """items: list of dicts with sample_id, stem, A..E, key. Writes one sheet per rule plus 'Start here' and 'Progress'."""
    from openpyxl.utils import get_column_letter
    blue = PatternFill("solid", fgColor="DDEBF7"); grey = PatternFill("solid", fgColor="EDEDED"); head = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="BFBFBF")
    wb = Workbook(); ws0 = wb.active; ws0.title = "Start here"
    helpers = {it["sample_id"]: helpers_for(it["stem"], {L: it[L] for L in "ABCDE"}, it["key"]) for it in items}
    for sheet, label, fields, helper, rule in BYRULE:
        ws = wb.create_sheet(sheet)
        cols = ["sample_id"]
        if "question" in fields: cols.append("question sentence (last sentence of the stem)")
        if "stem" in fields: cols.append("stem")
        if "options" in fields: cols += ["A", "B", "C", "D", "E"]
        if "key" in fields: cols.append("key")
        if helper: cols.append(helper)
        cols += ["ANSWER (1 = yes, 0 = no)", "notes"]
        ws.append([rule]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top"); ws.cell(row=1, column=1).font = Font(name="Arial", size=10, bold=True)
        ws.row_dimensions[1].height = 62
        ws.append(cols)
        for it in items:
            row = [it["sample_id"]]
            if "question" in fields: row.append(question_sentence(it["stem"]))
            if "stem" in fields: row.append(it["stem"])
            if "options" in fields: row += [it[L] for L in "ABCDE"]
            if "key" in fields: row.append(it["key"])
            if helper: row.append(helpers[it["sample_id"]][helper])
            row += [None, None]
            ws.append(row)
        n = len(items); ans_col = len(cols) - 1
        for j, c in enumerate(ws[2], 1):
            c.font = Font(bold=True, name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = blue if j == ans_col else (grey if cols[j - 1] == helper else head)
        for row in ws.iter_rows(min_row=3):
            for j, c in enumerate(row, 1):
                c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10)
                if j == ans_col: c.fill = blue
                elif cols[j - 1] == helper: c.fill = grey
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True, errorTitle="0 or 1 only", error="Type 1 for yes or 0 for no.")
        ws.add_data_validation(dv); dv.add("%s3:%s%d" % (get_column_letter(ans_col), get_column_letter(ans_col), n + 2))
        for j, name in enumerate(cols, 1):
            w = {"sample_id": 8, "stem": 60, "question sentence (last sentence of the stem)": 45, "key": 5, "ANSWER (1 = yes, 0 = no)": 12, "notes": 16}.get(name, 22)
            if name == helper: w = 48
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "B3"
    # Start here
    ws0.column_dimensions["A"].width = 120
    lines = [("Hand-check, one rule at a time (%s)" % pass_name, True), ("", False),
             ("There are ten sheets, one per rule, each listing the same %d questions. Take ONE sheet at a time, read the rule in its top row, and go down the blue ANSWER column typing 1 (yes) or 0 (no) for every row. Finish the sheet before you start the next one. Never leave an answer blank." % len(items), False), ("", False),
             ("Each sheet shows only what that rule needs. The key letter is shown only on sheets 4, 7 and 8, where the rule is about the key; on the other sheets you judge the options without knowing which is correct, which is what the rule wants. Grey cells are counts and word matches computed for you from the text; they never show the checker's verdict, and the 1/0 is always yours.", False), ("", False),
             ("Suggested order (quickest first): 2 meta_option, 3 combination, 1 negative_stem, 6 vague, 5 absolute, 10 overlapping, 9 nonparallel, 4 longest, 8 grammatical, 7 clang. Sheets 1, 2, 3 and 6 are fast scans; 7 clang takes the most thought.", False), ("", False),
             ("Notes column: 'category mix' (sheet 9), 'logical overlap' (sheet 10), or anything odd (no key, six options). Put the date in the notes of the first row you do in each sitting.", False), ("", False),
             ("Before you start: do the ten practice items in docs/MedIWF_Rater_Guide_v3.docx and check them against its answer key. While labelling, do not look at any detector output and do not ask an AI. The Progress sheet counts how many rows of each sheet are done. When all ten sheets show %d of %d, tell me and I will merge the sheets and run the analysis." % (len(items), len(items)), False)]
    for text, bold in lines:
        ws0.append([text]); c = ws0.cell(row=ws0.max_row, column=1); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=11, bold=bold)
    wp = wb.create_sheet("Progress"); wp.append(["sheet", "answered", "of"]); wp.column_dimensions["A"].width = 22
    for sheet, label, fields, helper, rule in BYRULE:
        ncols = 1 + ("question" in fields) + ("stem" in fields) + 5 * ("options" in fields) + ("key" in fields) + (1 if helper else 0) + 2
        col = get_column_letter(ncols - 1)
        wp.append([sheet, "=COUNT('%s'!%s3:%s%d)" % (sheet, col, col, len(items) + 2), len(items)])
    wp.append(["all sheets", "=SUM(B2:B11)", 10 * len(items)])
    for row in wp.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10, bold=(row[0].row == 1))
    wb.save(out_path); print("wrote", out_path, "(%d items x 10 rule sheets)" % len(items))


def main():
    ws1 = load_workbook(SRC).active
    hdr = [c.value for c in ws1[1]]
    assert hdr[:8] == ["sample_id", "stem", "A", "B", "C", "D", "E", "key"] and hdr[8:18] == LABELS and hdr[18] == "notes", hdr
    wb = Workbook(); ws = wb.active; ws.title = "pass1"
    helper_names = list(helpers_for("Which of the following is a?", {L: "x y z w." for L in "ABCDE"}, "A").keys())
    ws.append(hdr[:19] + helper_names)
    blue = PatternFill("solid", fgColor="DDEBF7"); grey = PatternFill("solid", fgColor="EDEDED"); head = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="BFBFBF")
    n = 0
    for r in ws1.iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        stem = r[1]; options = {L: r[2 + i] for i, L in enumerate("ABCDE")}; key = str(r[7] or "").strip().upper()[:1]
        h = helpers_for(stem, options, key)
        ws.append(list(r[:8]) + [None] * 11 + [h[k] for k in helper_names]); n += 1
    last_col = ws.max_column
    for j, c in enumerate(ws[1], 1):
        c.font = Font(bold=True, name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = head if j <= 8 else (blue if j <= 18 else (head if j == 19 else grey))
    for row in ws.iter_rows(min_row=2):
        for j, c in enumerate(row, 1):
            c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10)
            if 9 <= j <= 18: c.fill = blue
            elif j >= 20: c.fill = grey
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True, errorTitle="0 or 1 only", error="Type 1 for yes or 0 for no.")
    ws.add_data_validation(dv); dv.add("I2:R%d" % (n + 1))
    widths = {"A": 8, "B": 60, "C": 20, "D": 20, "E": 20, "F": 20, "G": 20, "H": 5}
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    from openpyxl.utils import get_column_letter
    for j in range(9, 19): ws.column_dimensions[get_column_letter(j)].width = 11
    ws.column_dimensions["S"].width = 18
    for j in range(20, last_col + 1): ws.column_dimensions[get_column_letter(j)].width = 42
    ws.freeze_panes = "C2"
    # How-to sheet
    wh = wb.create_sheet("How to"); wh.column_dimensions["A"].width = 120
    for text, bold in HOWTO:
        wh.append([text]); c = wh.cell(row=wh.max_row, column=1); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=11, bold=bold)
    wc = wb.create_sheet("Cheat sheet")
    for row in CHEAT: wc.append(list(row))
    for col, w in zip("ABC", [22, 40, 90]): wc.column_dimensions[col].width = w
    for row in wc.iter_rows():
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(name="Arial", size=10, bold=(row[0].row == 1))
    wb.save(OUT); print("wrote", OUT, "with", n, "items and", len(helper_names), "helper columns")
    items = []
    for r in ws1.iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        items.append({"sample_id": int(r[0]), "stem": r[1], **{L: r[2 + i] for i, L in enumerate("ABCDE")}, "key": str(r[7] or "").strip().upper()[:1]})
    build_byrule(items, root / "outputs/verification_sample_pass1_byrule.xlsx", "pass 1")


if __name__ == "__main__":
    main()
