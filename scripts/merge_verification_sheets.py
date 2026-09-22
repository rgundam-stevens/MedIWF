"""Fold the rater's one-sheet-per-rule workbook back into the flat layout that analyze_verification.py reads.
Reads outputs/verification_sample_pass1_byrule.xlsx (or the pass-2 file), takes the ANSWER column of each of the ten rule
sheets, checks that every row is 0 or 1, and writes the ten labels (plus the notes, prefixed with the sheet name) into a
copy of the flat workbook (outputs/verification_sample_pass1_simple.xlsx) saved as outputs/verification_sample_pass1_merged.xlsx.
Usage: python3 scripts/merge_verification_sheets.py [--pass 2]
"""
import sys, pathlib, argparse, collections
from openpyxl import load_workbook
root = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(root / "scripts"))
from make_verification_helpers import BYRULE, LABELS


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--pass", dest="p", type=int, default=1); a = ap.parse_args()
    src = root / ("outputs/verification_sample_pass%d_byrule.xlsx" % a.p)
    flat = root / ("outputs/verification_sample_pass%d_simple.xlsx" % a.p)
    if not flat.exists(): flat = root / ("outputs/verification_sample_pass%d.xlsx" % a.p)
    out = root / ("outputs/verification_sample_pass%d_merged.xlsx" % a.p)
    wb = load_workbook(src, data_only=True)
    labels = collections.defaultdict(dict); notes = collections.defaultdict(list); problems = []
    for sheet, label, fields, helper, rule in BYRULE:
        ws = wb[sheet]
        hdr = [c.value for c in ws[2]]
        ai = hdr.index("ANSWER (1 = yes, 0 = no)"); ni = hdr.index("notes")
        for row in ws.iter_rows(min_row=3, values_only=True):
            if row[0] is None: continue
            sid = int(row[0]); v = row[ai]
            if v is None or str(v).strip() == "":
                problems.append("%s: row %d is blank" % (sheet, sid)); labels[sid][label] = None; continue
            try:
                iv = int(float(str(v).strip()))
                if iv not in (0, 1): raise ValueError
            except ValueError:
                problems.append("%s: row %d has %r, not 0/1" % (sheet, sid, v)); labels[sid][label] = None; continue
            labels[sid][label] = iv
            if row[ni] not in (None, ""): notes[sid].append("%s: %s" % (sheet, str(row[ni]).strip()))
    fw = load_workbook(flat); ws = fw["pass%d" % a.p] if ("pass%d" % a.p) in fw.sheetnames else fw.active
    hdr = [c.value for c in ws[1]]
    col = {name: hdr.index(name) + 1 for name in LABELS}; ncol = hdr.index("notes") + 1
    n = 0
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(row=r, column=1).value
        if sid is None: continue
        sid = int(sid); n += 1
        for name in LABELS: ws.cell(row=r, column=col[name]).value = labels.get(sid, {}).get(name)
        if notes.get(sid): ws.cell(row=r, column=ncol).value = " | ".join(notes[sid])
    fw.save(out)
    done = sum(1 for sid in labels if all(labels[sid].get(name) is not None for name in LABELS))
    print("merged %d items into %s; fully labelled: %d; problems: %d" % (n, out.name, done, len(problems)))
    for pr in problems[:40]: print("  " + pr)
    if len(problems) > 40: print("  ... and %d more" % (len(problems) - 40))


if __name__ == "__main__":
    main()
